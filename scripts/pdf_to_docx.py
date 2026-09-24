#!/usr/bin/env python3
"""
Chuyển PDF -> DOCX có Heading 1/2 chuẩn để pipeline `sano-docx2tts` nạp được.

Quy ước map (khớp parser ở internal/bookmaker/docx.go):
  - Heading 1  -> chương (chapter)
  - Heading 2  -> tiểu mục (section)
  - Title      -> tiêu đề sách
  - Paragraph  -> nội dung gốc (GIỮ NGUYÊN lời sách, PRD §5 — không viết lại)

Chiến lược trích cấu trúc:
  1. Ưu tiên outline/bookmark của PDF:
       entry level 0 -> Heading 1, level >= 1 -> Heading 2.
       Văn bản mỗi entry = text các trang từ trang entry tới ngay trước entry kế.
       Mỗi trang chỉ gán cho 1 entry (tránh lặp text).
  2. Không có outline -> heuristic theo dòng:
       dòng khớp "Chương/Phần/Chapter ..."  -> Heading 1
       dòng khớp "Mục/Bài/Section ..."       -> Heading 2
       còn lại -> paragraph.
  3. Không bắt được heading nào -> fallback: 1 Heading 1 (tiêu đề) + toàn bộ
     text làm paragraph, để pipeline vẫn chạy được (sinh 1 chương 1 tiểu mục).

Usage:
  python scripts/pdf_to_docx.py <input.pdf> <output.docx>

Thư viện: pypdf, python-docx (xem scripts/requirements.txt). Chạy trong venv:
  python3 -m venv scripts/.venv-pdf2docx
  scripts/.venv-pdf2docx/bin/pip install -r scripts/requirements.txt
  scripts/.venv-pdf2docx/bin/python scripts/pdf_to_docx.py in.pdf out.docx
"""
import os
import re
import sys

try:
    from pypdf import PdfReader
    from docx import Document
except ImportError as exc:  # pragma: no cover - chỉ chạy khi thiếu lib
    sys.stderr.write(
        "lỗi: thiếu thư viện (%s). Cài: pip install -r scripts/requirements.txt\n" % exc
    )
    sys.exit(3)

# Marker tiếng Việt + Anh cho heuristic khi PDF không có outline.
CHAPTER_RE = re.compile(
    r"^\s*(chương|chuong|phần|phan|chapter|part)\s+[0-9IVXLCDM\.]+",
    re.IGNORECASE,
)
SECTION_RE = re.compile(
    r"^\s*(mục|muc|bài|bai|section)\s+[0-9\.]+",
    re.IGNORECASE,
)


# Giới hạn chống PDF độc (hàng chục nghìn trang / outline khổng lồ làm treo máy).
# Sách thật hiếm khi quá vài nghìn trang.
MAX_PAGES = 5000
MAX_OUTLINE_ENTRIES = 10000


def page_text(reader, idx):
    """Trích text 1 trang, trả chuỗi (rỗng nếu lỗi/không có)."""
    try:
        return reader.pages[idx].extract_text() or ""
    except Exception:  # pragma: no cover - trang lỗi không chặn cả file
        return ""


def add_body(doc, text):
    """Thêm nội dung gốc làm các paragraph, giữ nguyên lời (tách theo dòng)."""
    for line in text.splitlines():
        line = line.strip()
        if line:
            doc.add_paragraph(line)


def flatten_outline(reader):
    """Duyệt outline đệ quy -> list dict {level, title, page} theo thứ tự tài liệu."""
    flat = []

    def walk(items, level):
        for it in items:
            if isinstance(it, list):
                walk(it, level + 1)
                continue
            if len(flat) >= MAX_OUTLINE_ENTRIES:
                return
            title = str(getattr(it, "title", "") or "").strip()
            if not title:
                continue
            try:
                page = reader.get_destination_page_number(it)
            except Exception:
                page = None
            flat.append({"level": level, "title": title, "page": page})

    try:
        outline = reader.outline
    except Exception:
        outline = []
    if outline:
        walk(outline, 0)
    return flat


def convert_with_outline(reader, doc, flat):
    """Dựng docx từ outline: level 0 -> H1, sâu hơn -> H2. Mỗi trang gán 1 lần."""
    num_pages = len(reader.pages)
    emitted = set()
    for i, entry in enumerate(flat):
        level = 1 if entry["level"] == 0 else 2
        doc.add_heading(entry["title"], level=level)

        start = entry["page"] if entry["page"] is not None else 0
        end = num_pages
        if i + 1 < len(flat) and flat[i + 1]["page"] is not None:
            end = flat[i + 1]["page"]
        if end <= start:
            end = start + 1  # cùng trang: lấy ít nhất trang hiện tại
        for pg in range(start, min(end, num_pages)):
            if pg in emitted:
                continue
            emitted.add(pg)
            add_body(doc, page_text(reader, pg))


def convert_with_heuristic(reader, doc):
    """Không có outline: nhận heading theo marker dòng; còn lại là paragraph.

    Trả True nếu bắt được ít nhất 1 Heading 1.
    """
    has_h1 = False
    for idx in range(len(reader.pages)):
        for line in page_text(reader, idx).splitlines():
            stripped = line.strip()
            if not stripped:
                continue
            if CHAPTER_RE.match(stripped):
                doc.add_heading(stripped, level=1)
                has_h1 = True
            elif SECTION_RE.match(stripped):
                doc.add_heading(stripped, level=2)
            else:
                doc.add_paragraph(stripped)
    return has_h1


def book_title(reader, pdf_path):
    """Suy tiêu đề sách: metadata.title -> dòng đầu trang 1 -> tên file."""
    try:
        meta = reader.metadata
        if meta and meta.title and meta.title.strip():
            return meta.title.strip()
    except Exception:
        pass
    for line in page_text(reader, 0).splitlines():
        if line.strip():
            return line.strip()
    return os.path.splitext(os.path.basename(pdf_path))[0]


def main():
    if len(sys.argv) != 3:
        sys.stderr.write("usage: pdf_to_docx.py <input.pdf> <output.docx>\n")
        sys.exit(2)
    pdf_path, docx_path = sys.argv[1], sys.argv[2]
    if not os.path.isfile(pdf_path):
        sys.stderr.write("lỗi: không thấy file PDF %r\n" % pdf_path)
        sys.exit(2)

    reader = PdfReader(pdf_path)
    # PDF có mật khẩu hoặc khoá hạn chế (sao chép, in...) → không mở. Sano không có
    # tính năng gỡ khoá; chỉ dùng tài liệu bạn có quyền sử dụng.
    if reader.is_encrypted:
        sys.exit(
            "PDF có mật khẩu hoặc khoá bảo vệ, Sano không mở. "
            "Chỉ dùng tài liệu bạn có quyền sử dụng."
        )
    if len(reader.pages) > MAX_PAGES:
        sys.exit(
            "PDF có %d trang, vượt giới hạn %d trang. Tách file nhỏ hơn rồi chạy lại."
            % (len(reader.pages), MAX_PAGES)
        )
    doc = Document()

    flat = flatten_outline(reader)
    if flat:
        convert_with_outline(reader, doc, flat)
        mode = "outline (%d mục)" % len(flat)
    else:
        # Tiêu đề sách trước, rồi heuristic theo dòng.
        doc.add_heading(book_title(reader, pdf_path), level=0)  # level 0 = style Title
        got_h1 = convert_with_heuristic(reader, doc)
        if not got_h1:
            # Fallback cuối: gom toàn bộ text vào 1 chương để pipeline vẫn chạy.
            doc = Document()
            doc.add_heading(book_title(reader, pdf_path), level=0)
            doc.add_heading("Nội dung", level=1)
            for idx in range(len(reader.pages)):
                add_body(doc, page_text(reader, idx))
            mode = "fallback (1 chương)"
        else:
            mode = "heuristic theo dòng"

    doc.save(docx_path)
    print("✓ Đã ghi %s (chế độ: %s, %d trang PDF)" % (docx_path, mode, len(reader.pages)))


if __name__ == "__main__":
    main()
